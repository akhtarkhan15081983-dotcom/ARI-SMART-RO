import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsAdminOrManager, user_role
from .models import Customer
from .serializers import CustomerSerializer


MAX_IMPORT_ROWS = 5000
MAX_IMPORT_SIZE = 10 * 1024 * 1024


def customer_qr_payload(customer):
    return f"ARI-SMART-RO:CUSTOMER:{customer.customer_id}"


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _phone(value):
    digits = "".join(ch for ch in _clean(value) if ch.isdigit())
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    return digits[-10:]


def _decimal(value):
    text = _clean(value)
    if not text:
        return Decimal("0")
    try:
        return Decimal(text.replace(",", ""))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


ALIASES = {
    "name": {"name", "customer name", "custumer name", "customer_name"},
    "phone": {"phone", "mobile", "mobile no", "mobile number", "contact no", "contect no"},
    "alternate_phone": {"alternate phone", "alternate_phone", "alt phone"},
    "email": {"email", "email id"},
    "gender": {"gender"},
    "address": {"address", "adress"},
    "area": {"area"},
    "city": {"city"},
    "state": {"state"},
    "pincode": {"pincode", "pin code", "pin"},
    "ro_model": {"ro model", "ro_model", "model"},
    "installation_charge": {"installation charge", "instalation charge", "installation_charge"},
    "monthly_rent": {"monthly rent", "rent", "monthly_rent"},
    "security_deposit": {"security deposit", "security_deposit", "deposit"},
    "installation_date": {"installation date", "date of installation", "installation_date"},
    "old_card_number": {"old card number", "card number", "old_card_number"},
    "card_number": {"new card no.", "new card number", "card_number"},
}


def _normalized_header(value):
    return " ".join(_clean(value).lower().replace("_", " ").split())


def _map_headers(headers):
    mapped = {}
    for index, value in enumerate(headers):
        header = _normalized_header(value)
        for field, aliases in ALIASES.items():
            if header in aliases and field not in mapped:
                mapped[field] = index
                break
    return mapped


def _find_header_row(rows):
    for row_index, row in enumerate(rows[:5]):
        mapped = _map_headers(row)
        if "name" in mapped and "phone" in mapped:
            return row_index, mapped
    return None, {}


def _xlsx_rows(upload):
    upload.seek(0)
    workbook = load_workbook(upload, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def _csv_rows(upload):
    upload.seek(0)
    raw = upload.read()
    if isinstance(raw, bytes):
        text = raw.decode("utf-8-sig", errors="replace")
    else:
        text = raw
    return [row for row in csv.reader(io.StringIO(text))]


class CustomerBulkImportAPIView(APIView):
    permission_classes = [IsAdminOrManager]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        upload = request.FILES.get("file")
        if upload is None:
            return Response({"detail": "Excel or CSV file is required."}, status=status.HTTP_400_BAD_REQUEST)

        if upload.size > MAX_IMPORT_SIZE:
            return Response({"detail": "File is too large. Maximum size is 10 MB."}, status=status.HTTP_400_BAD_REQUEST)

        filename = (upload.name or "").lower()
        try:
            if filename.endswith(".xlsx"):
                rows = _xlsx_rows(upload)
            elif filename.endswith(".csv"):
                rows = _csv_rows(upload)
            else:
                return Response({"detail": "Only .xlsx and .csv files are supported."}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            return Response({"detail": f"Unable to read file: {exc}"}, status=status.HTTP_400_BAD_REQUEST)

        header_index, columns = _find_header_row(rows)
        if header_index is None:
            return Response(
                {"detail": "Header row not found. The file must contain at least Name and Phone columns."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        data_rows = rows[header_index + 1:]
        if len(data_rows) > MAX_IMPORT_ROWS:
            return Response(
                {"detail": f"Maximum {MAX_IMPORT_ROWS} customer rows can be imported at one time."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        preview_only = str(request.data.get("preview_only", "")).lower() in {"1", "true", "yes"}
        created = []
        errors = []
        duplicates = []
        seen_phones = set()

        for row_number, row in enumerate(data_rows, start=header_index + 2):
            def value(field):
                index = columns.get(field)
                return row[index] if index is not None and index < len(row) else None

            name = _clean(value("name"))
            phone = _phone(value("phone"))

            if not name and not phone:
                continue
            if not name:
                errors.append({"row": row_number, "error": "Customer name is required."})
                continue
            if len(phone) != 10 or phone[0] not in "6789":
                errors.append({"row": row_number, "name": name, "error": "Valid 10-digit mobile number is required."})
                continue
            if phone in seen_phones:
                duplicates.append({"row": row_number, "name": name, "phone": phone, "reason": "Duplicate phone inside file"})
                continue
            seen_phones.add(phone)

            existing = Customer.objects.filter(phone=phone).only("id", "customer_id", "name", "phone").first()
            if existing:
                duplicates.append({
                    "row": row_number,
                    "name": name,
                    "phone": phone,
                    "customer_id": existing.customer_id,
                    "reason": "Phone already exists",
                })
                continue

            payload = {
                "name": name,
                "phone": phone,
                "alternate_phone": _phone(value("alternate_phone")),
                "email": _clean(value("email")),
                "gender": _clean(value("gender")).upper(),
                "address": _clean(value("address")),
                "area": _clean(value("area")),
                "city": _clean(value("city")) or "Agra",
                "state": _clean(value("state")) or "Uttar Pradesh",
                "pincode": _clean(value("pincode")) or "282001",
                "ro_model": _clean(value("ro_model")) or "UNKNOWN",
                "installation_charge": _decimal(value("installation_charge")),
                "monthly_rent": _decimal(value("monthly_rent")),
                "security_deposit": _decimal(value("security_deposit")),
                "installation_date": _date(value("installation_date")),
                "old_card_number": _clean(value("old_card_number")),
                "card_number": _clean(value("card_number")),
                "is_active": True,
            }
            if payload["gender"] not in {"MALE", "FEMALE", "OTHER"}:
                payload["gender"] = ""

            if preview_only:
                created.append({
                    "row": row_number,
                    "name": name,
                    "phone": phone,
                    "status": "ready",
                })
                continue

            try:
                with transaction.atomic():
                    customer = Customer.objects.create(**payload)
            except Exception as exc:
                errors.append({"row": row_number, "name": name, "phone": phone, "error": str(exc)})
                continue

            created.append({
                "id": customer.id,
                "customer_id": customer.customer_id,
                "name": customer.name,
                "phone": customer.phone,
                "qr_payload": customer_qr_payload(customer),
            })

        return Response(
            {
                "success": len(errors) == 0,
                "preview_only": preview_only,
                "summary": {
                    "total_rows": len(data_rows),
                    "ready_or_created": len(created),
                    "duplicates": len(duplicates),
                    "errors": len(errors),
                },
                "customers": created,
                "duplicates": duplicates[:200],
                "errors": errors[:200],
            },
            status=status.HTTP_200_OK,
        )


class CustomerQRCodeAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        customer = Customer.objects.filter(pk=pk).first()
        if customer is None:
            return Response({"detail": "Customer not found."}, status=status.HTTP_404_NOT_FOUND)

        role = user_role(request.user)
        allowed = role in {"ADMIN", "MANAGER", "OFFICE"}
        if role == "ENGINEER":
            allowed = customer.assigned_engineer_id == getattr(getattr(request.user, "employee_profile", None), "id", None)
        elif role == "CUSTOMER":
            allowed = False
            if request.user.is_verified:
                if customer.user_id == request.user.id:
                    allowed = True
                elif customer.user_id is None and customer.phone == request.user.phone:
                    first_match = (
                        Customer.objects
                        .filter(phone=request.user.phone, user__isnull=True)
                        .order_by("id")
                        .only("id")
                        .first()
                    )
                    allowed = first_match is not None and first_match.id == customer.id

        if not allowed:
            return Response({"detail": "You do not have access to this customer QR."}, status=status.HTTP_403_FORBIDDEN)

        return Response({
            "customer": CustomerSerializer(customer, context={"request": request}).data,
            "qr_payload": customer_qr_payload(customer),
        })
